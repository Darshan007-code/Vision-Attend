"""
VisionAttend - Automated Face Recognition Attendance Tracking System
Export Service (Excel .xlsx, CSV, and ReportLab PDF Report Generation)
"""

import csv
import io
import sys
from datetime import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

import config

class ExportService:
    def export_csv(self, records: list) -> str:
        """Exports attendance records to a CSV string."""
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Header
        writer.writerow([
            "Record ID", "Roll Number", "Student Name", "Department", "Semester",
            "Subject Code", "Subject Name", "Date", "Time", "Status",
            "Confidence (%)", "Liveness Verified", "Session Type", "Verification Method"
        ])

        for r in records:
            writer.writerow([
                r['id'],
                r['roll_number'],
                r['student_name'],
                r['student_dept'],
                r['semester'],
                r['subject_code'],
                r['subject_name'],
                r['date'],
                r['time'],
                r['status'],
                f"{r['confidence']}%",
                "Yes" if r['liveness_verified'] else "No",
                r['session_type'],
                r['method']
            ])

        return output.getvalue()

    def export_excel(self, records: list, subject_name: str = "All Courses", date_str: str = "") -> str:
        """Generates a styled Excel (.xlsx) file and returns its path."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance Sheet"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
        sub_font = Font(name="Calibri", size=11, italic=True, color="475569")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
        
        present_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        present_font = Font(name="Calibri", size=10, bold=True, color="166534")
        late_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
        late_font = Font(name="Calibri", size=10, bold=True, color="92400E")

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        # Title Block
        ws.merge_cells("A1:K1")
        ws["A1"] = "VISIONATTEND - AUTOMATED ATTENDANCE REPORT"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        ws.merge_cells("A2:K2")
        ws["A2"] = f"Subject: {subject_name}  |  Date: {date_str or 'All Dates'}  |  Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws["A2"].font = sub_font
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        # Column Headers
        headers = [
            "#", "Roll No", "Student Name", "Department", "Semester",
            "Subject", "Date", "Time", "Status", "Confidence", "Method"
        ]
        
        ws.append([]) # Blank row 3
        ws.append(headers) # Row 4
        ws.row_dimensions[4].height = 26

        for col_num, _ in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Populate Data
        for idx, r in enumerate(records, 1):
            row_vals = [
                idx,
                r['roll_number'],
                r['student_name'],
                r['student_dept'],
                r['semester'],
                f"{r['subject_code']} - {r['subject_name']}",
                r['date'],
                r['time'],
                r['status'],
                f"{r['confidence']}%",
                r['method']
            ]
            ws.append(row_vals)
            row_idx = 4 + idx
            ws.row_dimensions[row_idx].height = 20

            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.border = thin_border
                c.alignment = Alignment(vertical="center", horizontal="center" if col_idx in [1, 2, 5, 7, 8, 9, 10] else "left")

                if col_idx == 9: # Status
                    if r['status'] == "Present":
                        c.fill = present_fill
                        c.font = present_font
                    elif r['status'] == "Late":
                        c.fill = late_fill
                        c.font = late_font

        # Auto-fit column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Save file to reports/
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"Attendance_Report_{timestamp}.xlsx"
        filepath = config.REPORTS_DIR / filename
        wb.save(str(filepath))
        return str(filepath)

    def export_pdf(self, records: list, subject_name: str = "All Courses", date_str: str = "") -> str:
        """Generates a professional academic PDF attendance report via ReportLab."""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"Attendance_Sheet_{timestamp}.pdf"
        filepath = config.REPORTS_DIR / filename

        doc = SimpleDocTemplate(
            str(filepath),
            pagesize=A4,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=16,
            leading=20,
            textColor=colors.HexColor('#0f172a'),
            alignment=1 # Center
        )
        subtitle_style = ParagraphStyle(
            'ReportSubtitle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#475569'),
            alignment=1
        )
        meta_style = ParagraphStyle(
            'MetaStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=12,
            textColor=colors.HexColor('#1e293b')
        )

        elements = []

        # 1. Header Banner
        elements.append(Paragraph("<b>DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING</b>", title_style))
        elements.append(Paragraph("VisionAttend: Automated Biometric Attendance Tracking System", subtitle_style))
        elements.append(Spacer(1, 10))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#2563eb'), spaceAfter=12))

        # 2. Metadata Box
        total_count = len(records)
        present_count = sum(1 for r in records if r['status'] == 'Present')
        late_count = sum(1 for r in records if r['status'] == 'Late')
        
        meta_data = [
            [
                Paragraph(f"<b>Course / Subject:</b> {subject_name}", meta_style),
                Paragraph(f"<b>Date:</b> {date_str or 'All Sessions'}", meta_style)
            ],
            [
                Paragraph(f"<b>Total Logged:</b> {total_count} Students", meta_style),
                Paragraph(f"<b>Present:</b> {present_count}  |  <b>Late:</b> {late_count}", meta_style)
            ],
            [
                Paragraph(f"<b>Report Generated:</b> {datetime.now().strftime('%d %b %Y, %I:%M %p')}", meta_style),
                Paragraph("<b>Verification Engine:</b> OpenCV LBPH + Anti-Spoofing", meta_style)
            ]
        ]
        meta_table = Table(meta_data, colWidths=[270, 250])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#e2e8f0')),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('LEFTPADDING', (0,0), (-1,-1), 8),
            ('RIGHTPADDING', (0,0), (-1,-1), 8),
        ]))
        elements.append(meta_table)
        elements.append(Spacer(1, 15))

        # 3. Main Attendance Table
        table_headers = ["#", "Roll No", "Student Name", "Department", "Time", "Status", "Confidence", "Signature"]
        table_data = [table_headers]

        for idx, r in enumerate(records, 1):
            table_data.append([
                str(idx),
                str(r['roll_number']),
                str(r['student_name']),
                str(r['student_dept']),
                str(r['time']),
                str(r['status']),
                f"{r['confidence']}%",
                "" # Blank for manual signature verification
            ])

        # Widths total: 520 (matches A4 printable width)
        rec_table = Table(table_data, colWidths=[24, 60, 130, 85, 55, 55, 55, 56])
        rec_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e40af')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('ALIGN', (2, 1), (2, -1), 'LEFT'), # Left-align student names
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')])
        ]))

        elements.append(rec_table)
        elements.append(Spacer(1, 35))

        # 4. Signatures Section
        sig_data = [
            ["___________________________", "___________________________"],
            ["Faculty In-Charge", "Head of Department (HOD)"],
            ["Date: ____________________", "Seal / Stamp"]
        ]
        sig_table = Table(sig_data, colWidths=[260, 260])
        sig_table.setStyle(TableStyle([
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#334155')),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        elements.append(sig_table)

        doc.build(elements)
        return str(filepath)

export_service = ExportService()
